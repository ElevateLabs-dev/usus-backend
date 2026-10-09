import asyncio
import csv
import io
import logging
import uuid
from datetime import datetime, timezone
from typing import Iterator, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from src.core.config import settings
from src.core.dependencies import UserRole
from src.core.security import (
    UNUSABLE_PASSWORD,
    generate_temporary_password,
    hash_password,
)
from src.domains.tenants.models import Tenant
from src.domains.users.models import User
from src.domains.users.schemas import MAX_BULK_TRAINEES, SkippedTrainee, TraineeCreate
from src.infrastructure.email.sender import send_email

logger = logging.getLogger("usus.users")

_EMAIL_HEADERS = {"email", "email address", "e-mail"}
_NAME_HEADERS = {"full_name", "full name", "name"}
TEMPLATE_HEADERS = ("email", "full_name")


async def list_tenant_users(
    db: AsyncSession, tenant_id: uuid.UUID, role: UserRole | None = None
) -> Sequence[User]:
    query = select(User).where(User.tenant_id == tenant_id)
    if role is not None:
        query = query.where(User.role == role.value)
    result = await db.execute(query.order_by(User.created_at.desc()))
    return result.scalars().all()


async def add_trainees(
    db: AsyncSession, tenant_id: uuid.UUID, entries: Sequence[TraineeCreate]
) -> tuple[list[User], list[SkippedTrainee]]:
    """
    Create trainee accounts for an organization in one transaction.

    Accounts start locked (no usable password); send_invites() sets a temporary
    password and emails it. Emails already registered, or repeated within the
    same batch, are skipped and reported back.
    """
    skipped: list[SkippedTrainee] = []
    unique: dict[str, TraineeCreate] = {}

    for entry in entries:
        email = entry.email.strip().lower()
        if email in unique:
            skipped.append(
                SkippedTrainee(email=email, reason="Duplicate in this upload")
            )
            continue
        unique[email] = entry

    if unique:
        result = await db.execute(
            select(func.lower(User.email)).where(
                func.lower(User.email).in_(list(unique))
            )
        )
        for existing in result.scalars().all():
            unique.pop(existing, None)
            skipped.append(
                SkippedTrainee(email=existing, reason="Email already registered")
            )

    users = [
        User(
            email=email,
            full_name=(entry.full_name or "").strip() or None,
            hashed_password=UNUSABLE_PASSWORD,
            role=UserRole.TRAINEE.value,
            tenant_id=tenant_id,
            is_active=True,
            must_change_password=True,
        )
        for email, entry in unique.items()
    ]

    if users:
        db.add_all(users)
        await db.commit()
        for user in users:
            await db.refresh(user)

    return users, skipped


def parse_trainee_file(
    filename: str, content: bytes
) -> tuple[list[TraineeCreate], list[SkippedTrainee]]:
    """
    Parse an uploaded trainee list: .xlsx (first sheet) or .csv.

    The first row must be headers with an `email` column; a `full_name` (or
    `name`) column is optional. Rows with an invalid email are returned as
    skipped. Raises ValueError for an unsupported/unreadable file, a missing
    email column, or more than MAX_BULK_TRAINEES rows.
    """
    name = filename.lower()
    if name.endswith(".xlsx") or content.startswith(b"PK"):
        rows = _xlsx_rows(content)
    elif name.endswith(".csv"):
        rows = _csv_rows(content)
    else:
        raise ValueError("Upload an Excel (.xlsx) or CSV (.csv) file")

    header = next(rows, None)
    if header is None:
        raise ValueError("The file is empty")
    columns = [str(h or "").strip().lower() for h in header]
    email_idx = next((i for i, h in enumerate(columns) if h in _EMAIL_HEADERS), None)
    name_idx = next((i for i, h in enumerate(columns) if h in _NAME_HEADERS), None)
    if email_idx is None:
        raise ValueError("The first row must contain an 'email' column header")

    def cell(row: Sequence, idx: int | None) -> str:
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]).strip()

    valid: list[TraineeCreate] = []
    skipped: list[SkippedTrainee] = []
    for row in rows:
        email = cell(row, email_idx)
        if not email:
            continue  # blank row
        if len(valid) + len(skipped) >= MAX_BULK_TRAINEES:
            raise ValueError(f"A file can contain at most {MAX_BULK_TRAINEES} trainees")
        try:
            valid.append(
                TraineeCreate(email=email, full_name=cell(row, name_idx) or None)
            )
        except ValidationError:
            skipped.append(SkippedTrainee(email=email, reason="Invalid email address"))

    return valid, skipped


def _xlsx_rows(content: bytes) -> Iterator[Sequence]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Could not read the Excel file; save it as .xlsx and retry")
    try:
        yield from workbook.worksheets[0].iter_rows(values_only=True)
    finally:
        workbook.close()


def _csv_rows(content: bytes) -> Iterator[Sequence]:
    try:
        text = content.decode("utf-8-sig")  # tolerate Excel's BOM
    except UnicodeDecodeError:
        raise ValueError("CSV must be UTF-8 encoded")
    yield from csv.reader(io.StringIO(text))


def build_import_template() -> bytes:
    """An .xlsx with the expected headers and one example row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trainees"
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(("jane.doe@yourcompany.com", "Jane Doe"))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 28
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _invite_email(user: User, organization: str, temporary_password: str) -> str:
    greeting = f"Hello {user.full_name}," if user.full_name else "Hello,"
    return (
        f"{greeting}\n\n"
        f"{organization} has added you to Usus, "
        "the customer-service training platform.\n\n"
        f"Sign in here: {settings.FRONTEND_URL.rstrip('/')}/login\n"
        f"Email: {user.email}\n"
        f"Temporary password: {temporary_password}\n\n"
        "You will be asked to choose a new password the first time you sign in.\n"
    )


async def send_invites(
    user_ids: Sequence[uuid.UUID], session_factory: async_sessionmaker[AsyncSession]
) -> None:
    """
    Background task: give each user a new temporary password and email it.

    Also used to resend an invite, which replaces any previous password.
    A failed email is logged and invite_sent_at stays unchanged, so the
    organization can see who still needs a resend.
    """
    async with session_factory() as db:
        for user_id in user_ids:
            user = await db.get(User, user_id)
            if user is None:
                continue

            tenant = await db.get(Tenant, user.tenant_id) if user.tenant_id else None
            organization = tenant.name if tenant else "Your organization"

            temporary_password = generate_temporary_password()
            # bcrypt is CPU-heavy (~0.4s); keep it off the event loop.
            user.hashed_password = await asyncio.to_thread(
                hash_password, temporary_password
            )
            user.must_change_password = True
            await db.commit()

            try:
                await asyncio.to_thread(
                    send_email,
                    user.email,
                    "You've been invited to Usus",
                    _invite_email(user, organization, temporary_password),
                )
            except Exception:
                logger.exception("Invite email to %s failed", user.email)
                continue

            user.invite_sent_at = datetime.now(timezone.utc)
            await db.commit()
