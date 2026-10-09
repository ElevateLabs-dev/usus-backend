"""
Organization-managed trainee onboarding:
add (single / bulk JSON / Excel or CSV import) -> invite email -> temporary-password sign-in
-> forced password change.

Runs on in-memory SQLite; emails are captured instead of sent.
"""

import io
import re
import uuid

import pytest
import pytest_asyncio
import httpx
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

import src.domains.evaluations.models  # noqa: F401  (register all tables)
import src.domains.scenarios.models  # noqa: F401
import src.domains.simulations.models  # noqa: F401
from src.core.base_model import Base
from src.core.config import settings
from src.core.database import get_db
from src.core.security import hash_password
from src.domains.tenants.models import Tenant
from src.domains.users import router as users_router
from src.domains.users import service as users_service
from src.domains.users.models import User
from src.infrastructure.email import sender
from src.main import app

ADMIN_PASSWORD = "admin-pass-123"


@pytest_asyncio.fixture
async def session_factory():
    # StaticPool: one shared in-memory DB for requests AND the background invite task.
    engine = create_async_engine(
        "sqlite+aiosqlite:///:memory:",
        poolclass=StaticPool,
        connect_args={"check_same_thread": False},
    )
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    yield factory
    await engine.dispose()


@pytest.fixture
def outbox(monkeypatch):
    sent: list[dict] = []

    def fake_send_email(to, subject, body):
        sent.append({"to": to, "subject": subject, "body": body})

    monkeypatch.setattr(users_service, "send_email", fake_send_email)
    return sent


@pytest_asyncio.fixture
async def client(session_factory, monkeypatch, outbox):
    async def override_get_db():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(users_router, "_session_factory", session_factory)
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as c:
        yield c
    app.dependency_overrides.pop(get_db, None)


async def _make_org(session_factory, name="Acme", admin_email="admin@acme.com"):
    async with session_factory() as db:
        tenant = Tenant(name=name)
        db.add(tenant)
        await db.flush()
        db.add(
            User(
                email=admin_email,
                hashed_password=hash_password(ADMIN_PASSWORD),
                role="company_admin",
                tenant_id=tenant.id,
            )
        )
        await db.commit()
        return tenant.id


async def _login(client, email, password):
    return await client.post(
        "/api/v1/auth/token", data={"username": email, "password": password}
    )


async def _auth(client, email, password):
    response = await _login(client, email, password)
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _temp_password(email_body: str) -> str:
    return re.search(r"Temporary password: (\S+)", email_body).group(1)


async def test_trainee_cannot_add_trainees(client, session_factory, outbox):
    await _make_org(session_factory)
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)
    await client.post(
        "/api/v1/users/trainees", json={"email": "t1@acme.com"}, headers=admin
    )
    temp = _temp_password(outbox[0]["body"])
    change = await client.post(
        "/api/v1/auth/change-password",
        json={"current_password": temp, "new_password": "trainee-pass-1"},
        headers=await _auth(client, "t1@acme.com", temp),
    )
    trainee = {"Authorization": f"Bearer {change.json()['access_token']}"}

    response = await client.post(
        "/api/v1/users/trainees", json={"email": "t2@acme.com"}, headers=trainee
    )
    assert response.status_code == 403


async def test_add_trainee_invite_and_forced_password_change(
    client, session_factory, outbox
):
    await _make_org(session_factory)
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)

    response = await client.post(
        "/api/v1/users/trainees",
        json={"email": "Jane@Acme.com", "full_name": "Jane Doe"},
        headers=admin,
    )
    assert response.status_code == 201, response.text
    assert response.json()["email"] == "jane@acme.com"
    assert response.json()["role"] == "trainee"
    assert response.json()["must_change_password"] is True

    # Invite email went out with a temporary password and the org name.
    assert len(outbox) == 1
    assert outbox[0]["to"] == "jane@acme.com"
    assert "Acme" in outbox[0]["body"] and "Hello Jane Doe" in outbox[0]["body"]
    temp = _temp_password(outbox[0]["body"])

    # First sign-in with the temporary password (email is case-insensitive).
    login = await _login(client, "JANE@acme.com", temp)
    assert login.status_code == 200
    assert login.json()["must_change_password"] is True
    pending = {"Authorization": f"Bearer {login.json()['access_token']}"}

    # Everything except /users/me and change-password is blocked.
    blocked = await client.get("/api/v1/scenarios/", headers=pending)
    assert blocked.status_code == 403
    assert blocked.json()["detail"] == "Password change required"
    me = await client.get("/api/v1/users/me", headers=pending)
    assert me.status_code == 200 and me.json()["invite_sent_at"] is not None

    wrong = await client.post(
        "/api/v1/auth/change-password",
        json={"current_password": "nope-nope", "new_password": "brand-new-pass"},
        headers=pending,
    )
    assert wrong.status_code == 400

    changed = await client.post(
        "/api/v1/auth/change-password",
        json={"current_password": temp, "new_password": "brand-new-pass"},
        headers=pending,
    )
    assert changed.status_code == 200
    assert changed.json()["must_change_password"] is False
    ready = {"Authorization": f"Bearer {changed.json()['access_token']}"}

    assert (await client.get("/api/v1/scenarios/", headers=ready)).status_code == 200
    assert (await _login(client, "jane@acme.com", temp)).status_code == 401
    assert (await _login(client, "jane@acme.com", "brand-new-pass")).json()[
        "must_change_password"
    ] is False


async def test_bulk_json_skips_duplicates_and_existing(client, session_factory, outbox):
    await _make_org(session_factory)
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)

    response = await client.post(
        "/api/v1/users/trainees/bulk",
        json={
            "trainees": [
                {"email": "a@acme.com", "full_name": "A"},
                {"email": "b@acme.com"},
                {"email": "A@acme.com"},  # duplicate (case-insensitive)
                {"email": "admin@acme.com"},  # already registered
            ]
        },
        headers=admin,
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert sorted(u["email"] for u in body["created"]) == ["a@acme.com", "b@acme.com"]
    assert {s["reason"] for s in body["skipped"]} == {
        "Duplicate in this upload",
        "Email already registered",
    }
    assert sorted(m["to"] for m in outbox) == ["a@acme.com", "b@acme.com"]

    listed = await client.get("/api/v1/users/?role=trainee", headers=admin)
    assert len(listed.json()) == 2


async def test_import_csv(client, session_factory, outbox):
    await _make_org(session_factory)
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)

    csv_content = (
        "Email,Full Name\nc1@acme.com,Chris One\nnot-an-email,Bad Row\n\nc2@acme.com,\n"
    ).encode()
    response = await client.post(
        "/api/v1/users/trainees/import",
        files={"file": ("trainees.csv", csv_content, "text/csv")},
        headers=admin,
    )
    assert response.status_code == 201, response.text
    body = response.json()
    created = {u["email"]: u["full_name"] for u in body["created"]}
    assert created == {"c1@acme.com": "Chris One", "c2@acme.com": None}
    assert body["skipped"] == [
        {"email": "not-an-email", "reason": "Invalid email address"}
    ]
    assert len(outbox) == 2

    no_email_col = await client.post(
        "/api/v1/users/trainees/import",
        files={"file": ("x.csv", b"name\nBob\n", "text/csv")},
        headers=admin,
    )
    assert no_email_col.status_code == 400


async def test_resend_invite_replaces_password_and_is_tenant_scoped(
    client, session_factory, outbox
):
    await _make_org(session_factory)
    await _make_org(session_factory, name="Other Co", admin_email="admin@other.com")
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)
    other_admin = await _auth(client, "admin@other.com", ADMIN_PASSWORD)

    created = await client.post(
        "/api/v1/users/trainees", json={"email": "r@acme.com"}, headers=admin
    )
    user_id = created.json()["id"]
    first_temp = _temp_password(outbox[0]["body"])

    other = await client.post(
        f"/api/v1/users/{user_id}/resend-invite", headers=other_admin
    )
    assert other.status_code == 404

    resent = await client.post(f"/api/v1/users/{user_id}/resend-invite", headers=admin)
    assert resent.status_code == 202
    second_temp = _temp_password(outbox[1]["body"])

    assert (await _login(client, "r@acme.com", first_temp)).status_code == 401
    assert (await _login(client, "r@acme.com", second_temp)).status_code == 200


async def test_failed_email_leaves_invite_unsent(client, session_factory, monkeypatch):
    await _make_org(session_factory)
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)

    def broken_send_email(to, subject, body):
        raise ConnectionError("Resend down")

    monkeypatch.setattr(users_service, "send_email", broken_send_email)
    response = await client.post(
        "/api/v1/users/trainees", json={"email": "f@acme.com"}, headers=admin
    )
    assert response.status_code == 201

    listed = await client.get("/api/v1/users/", headers=admin)
    trainee = next(u for u in listed.json() if u["email"] == "f@acme.com")
    assert trainee["invite_sent_at"] is None
    assert uuid.UUID(trainee["id"])


def _xlsx(rows) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def test_import_xlsx(client, session_factory, outbox):
    await _make_org(session_factory)
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)

    content = _xlsx(
        [
            ("Email Address", "Name"),
            ("x1@acme.com", "Xena One"),
            (None, None),  # blank row
            ("broken", "Bad"),
            ("x2@acme.com", None),
        ]
    )
    response = await client.post(
        "/api/v1/users/trainees/import",
        files={"file": ("staff.xlsx", content, "application/octet-stream")},
        headers=admin,
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert {u["email"]: u["full_name"] for u in body["created"]} == {
        "x1@acme.com": "Xena One",
        "x2@acme.com": None,
    }
    assert body["skipped"] == [{"email": "broken", "reason": "Invalid email address"}]
    assert sorted(m["to"] for m in outbox) == ["x1@acme.com", "x2@acme.com"]


async def test_import_template_round_trip(client, session_factory, outbox):
    await _make_org(session_factory)
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)

    template = await client.get("/api/v1/users/trainees/import-template", headers=admin)
    assert template.status_code == 200
    assert "attachment" in template.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(template.content)).active
    assert [c.value for c in sheet[1]] == ["email", "full_name"]

    # Fill the template the way an organization would, then upload it.
    sheet.delete_rows(2)
    sheet.append(("filled@acme.com", "Filled In"))
    buffer = io.BytesIO()
    sheet.parent.save(buffer)
    response = await client.post(
        "/api/v1/users/trainees/import",
        files={"file": ("template.xlsx", buffer.getvalue())},
        headers=admin,
    )
    assert [u["email"] for u in response.json()["created"]] == ["filled@acme.com"]


async def test_import_rejects_unsupported_files(client, session_factory):
    await _make_org(session_factory)
    admin = await _auth(client, "admin@acme.com", ADMIN_PASSWORD)

    response = await client.post(
        "/api/v1/users/trainees/import",
        files={"file": ("staff.pdf", b"%PDF-1.4 ...", "application/pdf")},
        headers=admin,
    )
    assert response.status_code == 400
    assert "xlsx" in response.json()["detail"]


def test_resend_sender_posts_email(monkeypatch):
    calls = []

    def fake_post(url, headers, json, timeout):
        calls.append({"url": url, "headers": headers, "json": json})
        return httpx.Response(200, json={"id": "email_123"})

    monkeypatch.setattr(settings, "RESEND_API_KEY", "re_test_key")
    monkeypatch.setattr(sender.httpx, "post", fake_post)
    sender.send_email("jane@acme.com", "Hi", "Body text")

    assert calls[0]["url"] == "https://api.resend.com/emails"
    assert calls[0]["headers"]["Authorization"] == "Bearer re_test_key"
    assert calls[0]["json"]["to"] == ["jane@acme.com"]
    assert calls[0]["json"]["text"] == "Body text"


def test_resend_sender_raises_on_rejection(monkeypatch):
    monkeypatch.setattr(settings, "RESEND_API_KEY", "re_test_key")
    monkeypatch.setattr(
        sender.httpx,
        "post",
        lambda *a, **k: httpx.Response(422, json={"message": "Invalid `from`"}),
    )
    with pytest.raises(RuntimeError, match="422"):
        sender.send_email("jane@acme.com", "Hi", "Body")


def test_sender_without_api_key_only_logs(monkeypatch, caplog):
    monkeypatch.setattr(settings, "RESEND_API_KEY", None)

    def must_not_post(*a, **k):
        raise AssertionError("should not call Resend without an API key")

    monkeypatch.setattr(sender.httpx, "post", must_not_post)
    sender.send_email("jane@acme.com", "Hi", "Body")
    assert "RESEND_API_KEY is not set" in caplog.text
